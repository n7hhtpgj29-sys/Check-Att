from __future__ import annotations
import os,re,sqlite3,threading,time,uuid
from datetime import datetime,date,time as dtime,timedelta
from pathlib import Path
from flask import Flask,jsonify,render_template,request,send_from_directory
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
BASE=Path(__file__).resolve().parent; DATA=BASE/'data'; DATA.mkdir(exist_ok=True); DB=DATA/'attendance.db'; TARGET='https://webapp.calcomp.co.th/att/'
app=Flask(__name__); app.config['MAX_CONTENT_LENGTH']=20*1024*1024; jobs={}; lock=threading.Lock()
def log(msg): print(f'[ATT] {datetime.now().isoformat(timespec="seconds")} {msg}', flush=True)
def db(): c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; return c
def init():
 with db() as c:
  c.execute('CREATE TABLE IF NOT EXISTS employees(employee_code TEXT PRIMARY KEY,full_name TEXT,department TEXT,position TEXT,active INTEGER DEFAULT 1,updated_at TEXT)')
  c.execute("CREATE TABLE IF NOT EXISTS attendance(employee_code TEXT PRIMARY KEY,name_from_web TEXT,latest_datetime TEXT,status TEXT,query_at TEXT,scan_in TEXT DEFAULT '',scan_out TEXT DEFAULT '',raw_count INTEGER DEFAULT 0,ot_minutes INTEGER DEFAULT 0,work_date TEXT DEFAULT '')")
  ec={x[1] for x in c.execute('PRAGMA table_info(employees)')}
  for n in ('location_support','team_support','group_code','shift'):
   if n not in ec:c.execute(f'ALTER TABLE employees ADD COLUMN {n} TEXT DEFAULT \'\'')
  ac={x[1] for x in c.execute('PRAGMA table_info(attendance)')}
  for n in ('scan_in','scan_out','work_date'):
   if n not in ac:c.execute(f"ALTER TABLE attendance ADD COLUMN {n} TEXT DEFAULT ''")
  if 'raw_count' not in ac:c.execute('ALTER TABLE attendance ADD COLUMN raw_count INTEGER DEFAULT 0')
  if 'ot_minutes' not in ac:c.execute('ALTER TABLE attendance ADD COLUMN ot_minutes INTEGER DEFAULT 0')
init()
def clean(v): return '' if v is None else str(int(v) if isinstance(v,float) and v.is_integer() else v).strip()
def normpos(v):
 s=clean(v); return {'op':'OP','staff':'Staff','engineer':'Engineer','supervisor':'Supervisor','manager':'Manager','vp':'VP'}.get(s.lower(),s.title() or 'Unknown')
def parse_dt(s):
 for f in ('%d/%m/%Y %H:%M:%S','%d/%m/%Y %H:%M'):
  try:return datetime.strptime(re.sub(r'\s+',' ',clean(s)),f)
  except:pass
 return None
def getrows(page):
 out=[]; rs=page.locator('table tbody tr'); rs=rs if rs.count() else page.locator('table tr')
 for i in range(rs.count()):
  cs=rs.nth(i).locator('td')
  if cs.count()>=4:out.append([clean(x) for x in cs.all_inner_texts()])
 return out
def workday_for_shift(shift, now=None, requested_work_date=None):
 if requested_work_date:
  try:return datetime.strptime(requested_work_date,'%Y-%m-%d').date()
  except:pass
 now=now or datetime.now()
 if clean(shift).upper()=='N' and now.time()<dtime(18,0): return now.date()-timedelta(days=1)
 return now.date()
def schedule(group,shift):
 g=clean(group).upper(); sh=clean(shift).upper()
 if sh=='N':
  return (dtime(19,40),dtime(4,40)) if g=='OP' else (dtime(20,0),dtime(5,40))
 return (dtime(7,40),dtime(16,40)) if g=='OP' else (dtime(8,0),dtime(17,40))
def choose(rs,emp,group='',shift='D',requested_work_date=None):
 valid=[]; name=''
 for r in rs:
  if len(r)<4 or clean(r[1]).upper()!=emp.upper():continue
  d=parse_dt(r[3]); name=name or clean(r[2])
  if d and d.year>=2020 and d.date()<=date.today()+timedelta(days=1):valid.append(d)
 if not valid:return None,'QUERY UNAVAILABLE',name,[],[],0,''
 wd=workday_for_shift(shift,requested_work_date=requested_work_date); sh=clean(shift).upper(); g=clean(group).upper(); start_t,end_t=schedule(g,sh)
 if sh=='N':
  window_start=datetime.combine(wd,dtime(16,0)); window_end=datetime.combine(wd+timedelta(days=1),dtime(12,0))
  rec=sorted(d for d in valid if window_start<=d<=window_end)
  ins=[d for d in rec if d.date()==wd and d.time()>=dtime(16,0)]
  outs=[d for d in rec if d.date()==wd+timedelta(days=1) and d.time()<=dtime(12,0)]
  scan_in=min(ins) if ins else None; scan_out=max(outs) if outs else None
  normal_end=datetime.combine(wd+timedelta(days=1),end_t)
 else:
  rec=sorted(d for d in valid if d.date()==wd)
  ins=[d for d in rec if d.time()<dtime(12,0)]
  outs=[d for d in rec if d.time()>=dtime(14,0)]
  scan_in=min(ins) if ins else None; scan_out=max(outs) if outs else None
  normal_end=datetime.combine(wd,end_t)
 if scan_in:
  # Staff who scan very early are treated as the automatic early-start pattern
  # (typically Monday / first day after a shutdown): 05:00 -> 14:00/14:40.
  # No schedule maintenance is required. OP keeps its own fixed schedule.
  if sh != 'N' and g != 'OP' and scan_in.time() < dtime(6,30):
   normal_end=datetime.combine(wd,dtime(14,40))
  # OT rule: subtract a 30-minute meal break, then require >= 60 net minutes.
  # Therefore a check mark starts only when scan-out is >= 90 minutes after normal end.
  gross=max(0,int((scan_out-normal_end).total_seconds()//60)) if scan_out else 0
  net=max(0,gross-30)
  ot=net if net>=60 else 0
  return max(rec), 'PRESENT', name, [scan_in], [scan_out] if scan_out else [], ot, wd.strftime('%d/%m/%Y')
 # Query succeeded but no scan in the selected shift window. Keep latest valid historical record.
 return max(valid),'NO SCAN TODAY',name,[],[],0,wd.strftime('%d/%m/%Y')
def query_once(page,emp,group="",shift="D",timeout=12,requested_work_date=None):
 page.goto(TARGET,wait_until='domcontentloaded',timeout=30000); inp=page.locator('input').first; inp.wait_for(state='visible',timeout=10000); inp.fill(emp)
 btn=page.get_by_role('button',name=re.compile(r'^\s*GO\s*$',re.I)); btn=btn if btn.count() else page.locator('button,input[type=submit]').first
 try:btn.click(no_wait_after=True,timeout=5000)
 except:pass
 end=time.time()+timeout; rs=[]
 while time.time()<end:
  try:
   rs=getrows(page)
   if any(len(r)>=4 and clean(r[1]).upper()==emp.upper() for r in rs):break
  except:pass
  page.wait_for_timeout(300)
 sel,status,name,ins,outs,ot_minutes,work_date=choose(rs,emp,group,shift,requested_work_date)
 if not sel:
  try:
   body=re.sub(r'\s+',' ',page.locator('body').inner_text())[:300]
   log(f'NO ROWS emp={emp} url={page.url} title={page.title()} rows={len(rs)} body={body}')
  except Exception as e: log(f'NO ROWS emp={emp} debug_error={type(e).__name__}: {e}')
  return {'latest_datetime':'','status':'QUERY UNAVAILABLE','name_from_web':'','scan_in':'','scan_out':'','raw_count':0,'ot_minutes':0,'work_date':''}
 return {'latest_datetime':sel.strftime('%d/%m/%Y %H:%M:%S'),'status':status,'name_from_web':name,'scan_in':ins[0].strftime('%d/%m/%Y %H:%M:%S') if ins else '','scan_out':outs[0].strftime('%d/%m/%Y %H:%M:%S') if outs else '','raw_count':len(ins)+len(outs),'ot_minutes':ot_minutes,'work_date':work_date}
def setjob(j,**kw):
 with lock:jobs[j].update(kw)
def runquery(j,query_shift='ALL',requested_work_date=None,missing_only=False):
 try:
  with db() as c:
   sql='SELECT * FROM employees WHERE active=1'; args=[]
   if query_shift in ('D','N'): sql+=' AND shift=?'; args.append(query_shift)
   # Incremental recheck: only employees not already PRESENT for the selected work date.
   # This preserves successful results and avoids re-querying the whole shift.
   if missing_only:
    target_wd=''
    if requested_work_date:
     try: target_wd=datetime.strptime(requested_work_date,'%Y-%m-%d').strftime('%d/%m/%Y')
     except: target_wd=''
    sql += " AND employee_code IN (SELECT e2.employee_code FROM employees e2 LEFT JOIN attendance a2 ON a2.employee_code=e2.employee_code WHERE e2.active=1 AND (a2.employee_code IS NULL OR a2.status IS NULL OR a2.status!='PRESENT' OR COALESCE(a2.work_date,'')!=?))"
    args.append(target_wd)
   sql+=' ORDER BY shift,location_support,team_support,employee_code'
   emps=[dict(x) for x in c.execute(sql,args)]
  setjob(j,status='running',total=len(emps),done=0,message='Checking attendance...'); results={}; failed=[]
  with sync_playwright() as p:
   headless=os.getenv('PLAYWRIGHT_HEADLESS','0')!='0'
   log(f'launch chromium headless={headless} display={os.getenv("DISPLAY","")} target={TARGET}')
   browser=p.chromium.launch(headless=headless, args=['--no-sandbox','--disable-dev-shm-usage','--disable-blink-features=AutomationControlled']); ctx=browser.new_context(ignore_https_errors=True, viewport={'width':1280,'height':900}, locale='en-US', timezone_id='Asia/Bangkok', user_agent='Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36'); ctx.add_init_script("Object.defineProperty(navigator,'webdriver',{get:()=>undefined});"); page=ctx.new_page()
   for i,e in enumerate(emps,1):
    emp=e['employee_code']; setjob(j,current=emp,done=i-1,message=f'Query {emp}')
    try:r=query_once(page,emp,e.get('group_code',''),e.get('shift','D'),requested_work_date=requested_work_date)
    except Exception as ex:
     log(f'QUERY ERROR emp={emp} {type(ex).__name__}: {ex}')
     r={'latest_datetime':'','status':'QUERY UNAVAILABLE','name_from_web':'','scan_in':'','scan_out':'','raw_count':0,'ot_minutes':0,'work_date':''}
    results[emp]=r
    if r['status']=='QUERY UNAVAILABLE':failed.append(emp)
    setjob(j,done=i); time.sleep(.6)
   if failed:
    for k,emp in enumerate(failed,1):
     setjob(j,current=emp,message=f'Retry {k}/{len(failed)}: {emp}')
     try:page.wait_for_timeout(1500); ee=next((x for x in emps if x['employee_code']==emp),{}); results[emp]=query_once(page,emp,ee.get('group_code',''),ee.get('shift','D'),15,requested_work_date)
     except Exception as ex: log(f'RETRY ERROR emp={emp} {type(ex).__name__}: {ex}')
   ctx.close(); browser.close()
  now=datetime.now().strftime('%d/%m/%Y %H:%M:%S')
  with db() as c:
   for emp,r in results.items():c.execute('INSERT INTO attendance(employee_code,name_from_web,latest_datetime,status,query_at,scan_in,scan_out,raw_count,ot_minutes,work_date) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(employee_code) DO UPDATE SET name_from_web=excluded.name_from_web,latest_datetime=excluded.latest_datetime,status=excluded.status,query_at=excluded.query_at,scan_in=excluded.scan_in,scan_out=excluded.scan_out,raw_count=excluded.raw_count,ot_minutes=excluded.ot_minutes,work_date=excluded.work_date',(emp,r['name_from_web'],r['latest_datetime'],r['status'],now,r.get('scan_in',''),r.get('scan_out',''),r.get('raw_count',0),r.get('ot_minutes',0),r.get('work_date','')))
  newly_present=sum(1 for r in results.values() if r.get('status')=='PRESENT')
  if missing_only:
   setjob(j,status='done',current='',done=len(emps),total=len(emps),message=f'Rechecked {len(emps)} employee(s) • {newly_present} present')
  else:
   setjob(j,status='done',current='',done=len(emps),total=len(emps),message='Attendance updated')
 except Exception as e:
  log(f'JOB ERROR {type(e).__name__}: {e}')
  setjob(j,status='error',current='',message=f'{type(e).__name__}: {e}')
@app.get('/manifest.webmanifest')
def manifest(): return send_from_directory(BASE/'static','manifest.webmanifest',mimetype='application/manifest+json')
@app.get('/sw.js')
def sw(): return send_from_directory(BASE/'static','sw.js',mimetype='application/javascript')
@app.get('/api/health')
def health(): return jsonify(ok=True,version='5.1-iphone-pwa-cloud-headed',target=TARGET,headless=os.getenv('PLAYWRIGHT_HEADLESS','0')!='0',tz=os.getenv('TZ',''))
@app.get('/')
def home():return render_template('index.html')
@app.get('/api/dashboard')
def dash():
 loc=request.args.get('location','ALL'); team=request.args.get('team','ALL'); sh=request.args.get('shift','ALL'); grp=request.args.get('group','ALL'); wh=['e.active=1']; args=[]
 if loc!='ALL':wh.append('e.location_support=?');args.append(loc)
 if team!='ALL':wh.append('e.team_support=?');args.append(team)
 if sh!='ALL':wh.append('e.shift=?');args.append(sh)
 if grp!='ALL':wh.append('e.group_code=?');args.append(grp)
 with db() as c:
  data=[dict(x) for x in c.execute(f"SELECT e.*,a.latest_datetime,a.scan_in,a.scan_out,a.status,a.query_at,a.ot_minutes,a.work_date FROM employees e LEFT JOIN attendance a ON a.employee_code=e.employee_code WHERE {' AND '.join(wh)} ORDER BY e.shift,e.location_support,e.team_support,e.employee_code",args)]
  locs=[x[0] for x in c.execute("SELECT DISTINCT location_support FROM employees WHERE active=1 AND COALESCE(location_support,'')<>'' ORDER BY location_support")]; teams=[x[0] for x in c.execute("SELECT DISTINCT team_support FROM employees WHERE active=1 AND COALESCE(team_support,'')<>'' ORDER BY team_support")]
 def stat(rows):
  out={'TOTAL':len(rows),'PRESENT':0,'NO SCAN TODAY':0,'QUERY UNAVAILABLE':0,'NOT CHECKED':0}
  for x in rows:
   st=x.get('status') or 'NOT CHECKED'; st=st if st in out else 'NOT CHECKED'; out[st]+=1
  return out
 def breakdown(rows,key):
  out={}
  for x in rows:
   k=clean(x.get(key)) or '(Not set)'; v=out.setdefault(k,{'total':0,'present':0,'no_scan':0,'error':0,'not_checked':0});v['total']+=1
   st=x.get('status') or 'NOT CHECKED'
   if st=='PRESENT':v['present']+=1
   elif st=='NO SCAN TODAY':v['no_scan']+=1
   elif st=='QUERY UNAVAILABLE':v['error']+=1
   else:v['not_checked']+=1
  return out
 day=[x for x in data if clean(x.get('shift')).upper()=='D']; night=[x for x in data if clean(x.get('shift')).upper()=='N']
 return jsonify(ok=True,employees=data,counts=stat(data),locations=locs,teams=teams,shift_counts={'D':stat(day),'N':stat(night)},by_location={'ALL':breakdown(data,'location_support'),'D':breakdown(day,'location_support'),'N':breakdown(night,'location_support')},by_team={'ALL':breakdown(data,'team_support'),'D':breakdown(day,'team_support'),'N':breakdown(night,'team_support')},last_query=max([x.get('query_at') or '' for x in data],default=''),current_night_work_date=workday_for_shift('N').strftime('%d/%m/%Y'),current_day_work_date=date.today().strftime('%d/%m/%Y'))
@app.post('/api/upload-master')
def upload():
 f=request.files.get('file')
 if not f or not f.filename.lower().endswith('.xlsx'):return jsonify(ok=False,error='กรุณาเลือกไฟล์ .xlsx'),400
 path=DATA/f'master_{uuid.uuid4().hex[:8]}.xlsx';f.save(path)
 try:
  wb=load_workbook(path,read_only=True,data_only=True);ws=wb[wb.sheetnames[0]]
  def hnorm(v):
   z=clean(v).lower().replace('\n',' ').replace('\r',' ');return re.sub(r'[^a-z0-9ก-๙]+','_',z).strip('_')
  hs=[hnorm(c.value) for c in ws[1]]
  aliases={'employee_code':['employee_code','employee_no','emp_no','empno','รหัสพนักงาน'],'full_name':['full_name','name','employee_name','ชื่อ','ชื่อพนักงาน'],'location_support':['location_support','location','support_location'],'team_support':['team_support','team','support_team'],'group_code':['group','group_code','group_support'],'shift':['shift','shift_code']}
  idx={}
  for k,names in aliases.items():
   for a in names:
    if hnorm(a) in hs:idx[k]=hs.index(hnorm(a));break
  req=['employee_code','full_name','location_support','team_support','group_code','shift'];missing=[k for k in req if k not in idx]
  if missing:raise ValueError('Excel ต้องมีคอลัมน์: employee_code, full_name, location_support, team_support, group, shift')
  incoming={}
  for r in ws.iter_rows(min_row=2,values_only=True):
   emp=re.sub(r'<[^>]+>','',clean(r[idx['employee_code']])).strip()
   if emp:incoming[emp]=(clean(r[idx['full_name']]),clean(r[idx['location_support']]),clean(r[idx['team_support']]),clean(r[idx['group_code']]).upper(),clean(r[idx['shift']]).upper())
  wb.close();now=datetime.now().isoformat(timespec='seconds')
  with db() as c:
   current={x[0] for x in c.execute('SELECT employee_code FROM employees WHERE active=1')}
   sql="INSERT INTO employees(employee_code,full_name,department,position,active,updated_at,location_support,team_support,group_code,shift) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(employee_code) DO UPDATE SET full_name=excluded.full_name,active=1,updated_at=excluded.updated_at,location_support=excluded.location_support,team_support=excluded.team_support,group_code=excluded.group_code,shift=excluded.shift"
   for emp,(name,loc,team,grp,sh) in incoming.items():c.execute(sql,(emp,name,'','',1,now,loc,team,grp,sh))
   c.executemany('UPDATE employees SET active=0,updated_at=? WHERE employee_code=?',[(now,x) for x in current-set(incoming)])
  return jsonify(ok=True,total=len(incoming),added=len(set(incoming)-current),removed=len(current-set(incoming)))
 except Exception as e:return jsonify(ok=False,error=str(e)),400
 finally:
  try:path.unlink(missing_ok=True)
  except:pass
@app.get('/api/employees')
def employee_list():
 with db() as c:return jsonify(ok=True,employees=[dict(x) for x in c.execute('SELECT * FROM employees ORDER BY active DESC,location_support,employee_code')])
@app.post('/api/employees')
def employee_save():
 x=request.get_json(force=True); emp=clean(x.get('employee_code'))
 if not emp:return jsonify(ok=False,error='Employee No. required'),400
 now=datetime.now().isoformat(timespec='seconds')
 with db() as c:c.execute('INSERT INTO employees(employee_code,full_name,department,position,active,updated_at,location_support,team_support,group_code,shift) VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(employee_code) DO UPDATE SET full_name=excluded.full_name,active=excluded.active,updated_at=excluded.updated_at,location_support=excluded.location_support,team_support=excluded.team_support,group_code=excluded.group_code,shift=excluded.shift',(emp,clean(x.get('full_name')),'','',1 if x.get('active',True) else 0,now,clean(x.get('location_support')),clean(x.get('team_support')),clean(x.get('group_code')).upper(),clean(x.get('shift')).upper()))
 return jsonify(ok=True)
@app.post('/api/query')
def start():
 x=request.get_json(silent=True) or {}; mode=clean(x.get('mode')).upper() or 'ALL'; today=date.today()
 if mode=='DAY': query_shift='D'; wd=today
 elif mode=='LAST_NIGHT': query_shift='N'; wd=today-timedelta(days=1)
 elif mode=='TONIGHT': query_shift='N'; wd=today
 else:
  query_shift=clean(x.get('shift')).upper() if clean(x.get('shift')).upper() in ('D','N') else 'ALL'
  try: wd=datetime.strptime(clean(x.get('work_date')),'%Y-%m-%d').date() if x.get('work_date') else None
  except: wd=None
 j=uuid.uuid4().hex[:10]
 missing_only=bool(x.get('missing_only',False))
 with lock:jobs[j]={'status':'queued','total':0,'done':0,'current':'','message':'Preparing missing employees...' if missing_only else 'Preparing...','query_shift':query_shift,'work_date':wd.isoformat() if wd else '','missing_only':missing_only}
 threading.Thread(target=runquery,args=(j,query_shift,wd.isoformat() if wd else None,missing_only),daemon=True).start();return jsonify(ok=True,job_id=j)
@app.get('/api/job/<j>')
def job(j):
 with lock:x=jobs.get(j)
 return jsonify(ok=bool(x),**(x or {'error':'Job not found'}))
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.getenv('PORT','5000')),debug=False,threaded=True)
